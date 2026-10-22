"""
Relatórios — análise de um período, sempre comparada com outro.

O Início mostra o mês de hoje e o Financeiro mostra um mês por vez. Aqui a organizadora escolhe um período
(mês, trimestre, ano ou datas livres) e vê como ele se saiu em relação ao período anterior ou ao mesmo período
do ano passado. Cada aba responde a uma pergunta do negócio:

- Resultado: quanto entrou, quanto custou e quanto sobrou (DRE simples, mês a mês)?
- Clientes: quem traz o dinheiro, de onde vêm e se voltam?
- Propostas: as propostas fecham? em quanto tempo? quais serviços fecham mais?
- Trabalhos: quais trabalhos deram lucro de verdade?
- Próximos meses: o que ainda vai entrar e sair, o que está atrasado e quais propostas aguardam resposta.

Boas práticas seguidas: números com definição visível, comparação com uma base (período anterior ou ano passado),
indicadores do que vem pela frente (a receber e propostas aguardando) em vez de projeções estatísticas, poucos
gráficos e exportação para Excel. Os cálculos ficam em utils/relatorios_calc.py (com testes).
"""
import html as html_module
import io
from datetime import datetime, timedelta, timezone

import pandas as pd
import plotly.graph_objects as go
import streamlit as st

from utils import relatorios_calc as rc
from utils.currency_formatter import fmt_brl
from utils.indicadores import cor_categoria

NAVY, GOLD, VERDE, VERMELHO, CINZA = "#0D1B2A", "#C9A84C", "#067647", "#B42318", "#98A2B3"

CSS = """
<style>
.rl-intro { color: #667085; font-size: .9rem; margin: -4px 0 14px; }
.rl-kpis { display: grid; grid-template-columns: repeat(4, minmax(0, 1fr)); gap: 12px; margin: 6px 0 14px; }
.rl-kpi { background: #fff; border: 1px solid #E4E7EC; border-radius: 12px; padding: 14px 16px; }
.rl-kpi .rot { font-size: .72rem; font-weight: 600; letter-spacing: .06em; text-transform: uppercase; color: #667085; }
.rl-kpi .val { font-size: 1.4rem; font-weight: 700; color: #101828; margin: 4px 0 2px; font-variant-numeric: tabular-nums; white-space: nowrap; }
.rl-kpi .sub { font-size: .8rem; color: #667085; line-height: 1.4; display: block; }
.rl-kpi.destaque { background: var(--navy, #0D1B2A); border-color: var(--navy, #0D1B2A); }
.rl-kpi.destaque .rot, .rl-kpi.destaque .sub { color: #C8D1DC; }
.rl-kpi.destaque .val { color: #fff; }
.rl-d { font-weight: 600; }
.rl-d.up { color: #067647; } .rl-d.down { color: #B42318; } .rl-d.neutro { color: #667085; }
.rl-kpi.destaque .rl-d.up { color: #6CE9A6; } .rl-kpi.destaque .rl-d.down { color: #FDA29B; }
.rl-secao { font-size: .78rem; font-weight: 600; letter-spacing: .06em; text-transform: uppercase; color: #667085; margin: 14px 0 8px; }
.rl-insight { background: #F8F9FB; border: 1px solid #E4E7EC; border-left: 3px solid var(--gold, #C9A84C); border-radius: 10px;
              padding: 10px 14px; font-size: .88rem; color: #344054; margin: 0 0 12px; line-height: 1.55; }
.rl-insight p { margin: 0 0 4px; } .rl-insight p:last-child { margin: 0; }
.rl-alerta { background: #FFFAEB; border: 1px solid #FEC84B; border-radius: 10px; padding: 10px 14px; font-size: .88rem; color: #93370D; margin: 0 0 12px; }
.rl-dre { width: 100%; border-collapse: collapse; font-size: .9rem; }
.rl-dre th { text-align: right; font-size: .72rem; font-weight: 600; letter-spacing: .06em; text-transform: uppercase; color: #667085;
             padding: 6px 10px; border-bottom: 1px solid #E4E7EC; }
.rl-dre th:first-child { text-align: left; }
.rl-dre td { padding: 8px 10px; border-bottom: 1px solid #F2F4F7; text-align: right; font-variant-numeric: tabular-nums; white-space: nowrap; }
.rl-dre td:first-child { text-align: left; color: #344054; }
.rl-dre-wrap { overflow-x: auto; }
.rl-dre tr.tot td { font-weight: 700; color: #101828; border-top: 1px solid #D0D5DD; }
.rl-dre tr.sub td:first-child { padding-left: 22px; color: #667085; }
.rl-cat { display: grid; grid-template-columns: minmax(0, 1.3fr) minmax(0, 1fr) auto; gap: 10px; align-items: center; padding: 5px 0; font-size: .86rem; }
.rl-cat .nm { display: flex; align-items: center; gap: 8px; color: #344054; min-width: 0; }
.rl-cat .nm span { overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
.rl-dot { flex: none; width: 10px; height: 10px; border-radius: 3px; display: inline-block; }
.rl-cat .bar { height: 8px; background: #F2F4F7; border-radius: 999px; overflow: hidden; }
.rl-cat .bar i { display: block; height: 100%; border-radius: 999px; }
.rl-cat b { font-variant-numeric: tabular-nums; white-space: nowrap; }
.rl-vazio { color: #98A2B3; font-size: .86rem; }
.rl-def { font-size: .78rem; color: #98A2B3; margin-top: 6px; line-height: 1.5; }
@media (max-width: 900px) { .rl-kpis { grid-template-columns: repeat(2, minmax(0, 1fr)); } .rl-kpi .val { font-size: 1.15rem; } }
</style>
"""


def _hoje():
    return datetime.now(timezone(timedelta(hours=-3))).date()


def _pct(v, casas=0):
    return "—" if v is None else f"{v * 100:.{casas}f}%".replace(".", ",")


def _delta(atual, comp, rotulo_comp, inverso=False, pontos=False):
    """'▲ 12% vs set/2026' com cor (inverso=True quando subir é ruim, ex.: gastos)."""
    if comp is None or atual is None:
        return f'<span class="rl-d neutro">sem base em {rotulo_comp}</span>'
    if pontos:   # diferença de margem em pontos percentuais
        d = (atual - comp) * 100
        if abs(d) < 0.05:
            return f'<span class="rl-d neutro">igual a {rotulo_comp}</span>'
        cls = "up" if (d > 0) != inverso else "down"
        return f'<span class="rl-d {cls}">{"▲" if d > 0 else "▼"} {abs(d):.1f} p.p.</span> vs {rotulo_comp}'.replace(".", ",", 1)
    v = rc.variacao(atual, comp)
    if v is None:
        return f'<span class="rl-d neutro">sem base em {rotulo_comp}</span>'
    if abs(v) < 0.005:
        return f'<span class="rl-d neutro">igual a {rotulo_comp}</span>'
    cls = "up" if (v > 0) != inverso else "down"
    return f'<span class="rl-d {cls}">{"▲" if v > 0 else "▼"} {_pct(abs(v))}</span> vs {rotulo_comp}'


def _kpi(rotulo, valor, sub, cls=""):
    return (f'<div class="rl-kpi {cls}"><span class="rot">{rotulo}</span>'
            f'<div class="val">{valor}</div><span class="sub">{sub}</span></div>')


def _kpis(*cards):
    st.markdown('<div class="rl-kpis">' + "".join(cards) + '</div>', unsafe_allow_html=True)


def _secao(txt):
    st.markdown(f'<div class="rl-secao">{txt}</div>', unsafe_allow_html=True)


def _insight(frases):
    if frases:
        st.markdown('<div class="rl-insight">' + "".join(f"<p>{html_module.escape(f)}</p>" for f in frases) + '</div>',
                    unsafe_allow_html=True)


def _barras(df):
    if df.empty:
        return '<div class="rl-vazio">Nada neste período.</div>'
    topo = float(df["valor"].iloc[0]) or 1
    h = ""
    for nome, valor in zip(df["categoria"].head(8), df["valor"].head(8)):
        cor, nome_h = cor_categoria(nome), html_module.escape(str(nome))
        h += (f'<div class="rl-cat"><span class="nm"><i class="rl-dot" style="background:{cor}"></i>'
              f'<span title="{nome_h}">{nome_h}</span></span>'
              f'<span class="bar"><i style="width:{valor / topo * 100:.0f}%;background:{cor}"></i></span>'
              f'<b>{fmt_brl(valor)}</b></div>')
    return h


def _layout(fig, altura=320):
    fig.update_layout(template="plotly_white", height=altura, margin=dict(l=8, r=8, t=10, b=8), separators=",.",
                      legend=dict(orientation="h", yanchor="bottom", y=1.02, x=0), hovermode="x unified",
                      font=dict(family="Inter, sans-serif", size=12, color="#344054"))
    fig.update_yaxes(gridcolor="#F2F4F7", zeroline=True, zerolinecolor="#D0D5DD", tickprefix="R$ ")
    return fig


def _tabela(df, moeda=(), pct=(), renomear=None, altura=None):
    d = df.copy()
    for c in moeda:
        if c in d:
            d[c] = d[c].apply(lambda v: "—" if v is None or pd.isna(v) else fmt_brl(v))
    for c in pct:
        if c in d:
            d[c] = d[c].apply(lambda v: _pct(v) if v is not None and not pd.isna(v) else "—")
    if renomear:
        d = d.rename(columns=renomear)[list(renomear.values())]
    kw = {"height": altura} if altura else {}
    st.dataframe(d, hide_index=True, use_container_width=True, **kw)


# ───────────────────────────── dados ─────────────────────────────

def _carregar():
    db = st.session_state.db
    fin = db.get_financeiro()
    # get_financeiro() traz no máximo 1.000 transações (e guarda em cache): para um ano inteiro pode faltar dado
    if fin is not None and len(fin) >= 1000:
        fin = db.get_financeiro(limit=200000, force_reload=True)
    try:
        propostas = db.get_propostas()
    except Exception:
        propostas = pd.DataFrame()
    try:
        clientes = db.get_clientes()
    except Exception:
        clientes = pd.DataFrame()
    fin = fin if fin is not None else pd.DataFrame()
    return rc.movimentos_rel(fin), propostas, clientes


# ───────────────────────────── abas ─────────────────────────────

def _aba_resultado(mov, ini, fim, cini, cfim, rot, rot_comp, exp):
    a, c = rc.resultado(mov, ini, fim), rc.resultado(mov, cini, cfim)
    cat_e, cat_s = rc.por_categoria(mov, ini, fim, "entra"), rc.por_categoria(mov, ini, fim, "sai")
    cat_s_comp = rc.por_categoria(mov, cini, cfim, "sai")
    if not a["faturamento"] and not (a["custos"] + a["despesas"]):
        st.info("Nenhum recebimento ou pagamento registrado neste período. Marque as receitas como Recebido e as despesas "
                "como Pago no Financeiro para elas aparecerem aqui.")
    _kpis(
        _kpi("Faturamento", fmt_brl(a["faturamento"]), _delta(a["faturamento"], c["faturamento"], rot_comp)),
        _kpi("Lucro", fmt_brl(a["lucro"]), _delta(a["lucro"], c["lucro"], rot_comp), "destaque"),
        _kpi("Margem de lucro", _pct(a["margem"], 1),
             _delta(a["margem"], c["margem"], rot_comp, pontos=True) if a["margem"] is not None and c["margem"] is not None
             else "lucro ÷ faturamento"),
        _kpi("Gastos da empresa", fmt_brl(a["custos"] + a["despesas"]),
             _delta(a["custos"] + a["despesas"], c["custos"] + c["despesas"], rot_comp, inverso=True)),
    )
    _insight(rc.destaques(a, c, cat_s, cat_s_comp, rot_comp))

    serie = rc.serie_mensal(mov, ini, fim)
    if len(serie) > 1:
        _secao("Mês a mês")
        fig = go.Figure()
        fig.add_bar(x=serie["mes"], y=serie["faturamento"], name="Faturamento", marker_color=NAVY,
                    hovertemplate="R$ %{y:,.2f}<extra>Faturamento</extra>")
        fig.add_bar(x=serie["mes"], y=serie["custos"] + serie["despesas"], name="Gastos", marker_color="#D0D5DD",
                    hovertemplate="R$ %{y:,.2f}<extra>Gastos</extra>")
        fig.add_scatter(x=serie["mes"], y=serie["lucro"], name="Lucro", mode="lines+markers",
                        line=dict(color=GOLD, width=3), hovertemplate="R$ %{y:,.2f}<extra>Lucro</extra>")
        fig.update_layout(barmode="group", bargap=0.35)
        st.plotly_chart(_layout(fig), use_container_width=True, config={"displayModeBar": False})

    # demonstrativo em largura total; categorias lado a lado embaixo
    if True:
        _secao("Demonstrativo do resultado")
        linhas = [("Faturamento", "faturamento", ""), ("(−) Custos diretos", "custos", "sub"),
                  ("= Margem bruta", "margem_bruta", "tot"), ("(−) Despesas fixas", "despesas", "sub"),
                  ("= Lucro", "lucro", "tot"), ("(−) Retiradas", "retiradas", "sub"),
                  ("= Ficou na empresa", "caixa", "tot")]
        h = (f'<div class="rl-dre-wrap"><table class="rl-dre"><tr><th></th><th>{html_module.escape(rot)}</th>'
             f'<th>{html_module.escape(rot_comp)}</th><th>Variação</th></tr>')
        for rot, k, cls in linhas:
            v = rc.variacao(a[k], c[k])
            h += (f'<tr class="{cls}"><td>{rot}</td><td>{fmt_brl(a[k])}</td><td>{fmt_brl(c[k])}</td>'
                  f'<td>{"—" if v is None else "0%" if abs(v) < 0.005 else ("+" if v > 0 else "−") + _pct(abs(v))}</td></tr>')
        h += (f'<tr><td>Margem de lucro</td><td>{_pct(a["margem"], 1)}</td><td>{_pct(c["margem"], 1)}</td><td></td></tr></table></div>'
              '<div class="rl-def">Regime de caixa: cada valor conta na data em que entrou ou saiu. Custos diretos são as '
              'despesas ligadas a um trabalho (assistente, equipe, fornecedor, comissão, material ou lançadas na proposta).</div>')
        st.markdown(h, unsafe_allow_html=True)
    col_e, col_s = st.columns(2, gap="large")
    with col_e:
        _secao("De onde veio")
        with st.container(border=True):
            st.markdown(_barras(cat_e), unsafe_allow_html=True)
    with col_s:
        _secao("Para onde foi")
        with st.container(border=True):
            st.markdown(_barras(cat_s), unsafe_allow_html=True)

    exp["Resultado mês a mês"] = serie.drop(columns=["ano", "m"]).rename(columns={
        "mes": "Mês", "faturamento": "Faturamento", "custos": "Custos diretos", "despesas": "Despesas",
        "lucro": "Lucro", "retiradas": "Retiradas"})
    exp["Demonstrativo"] = pd.DataFrame([{"Linha": r, rot: a[k], rot_comp: c[k]} for r, k, _ in linhas])
    exp["Receitas por categoria"] = cat_e.rename(columns={"categoria": "Categoria", "valor": "Valor"})
    exp["Gastos por categoria"] = cat_s.rename(columns={"categoria": "Categoria", "valor": "Valor"})


def _aba_clientes(propostas, clientes, ini, fim, cini, cfim, rot_comp, exp):
    a = rc.clientes_periodo(propostas, clientes, ini, fim)
    c = rc.clientes_periodo(propostas, clientes, cini, cfim)
    if not a["origem"].empty:   # no cadastro "Cliente" quer dizer indicação de outra cliente
        a["origem"] = a["origem"].assign(origem=a["origem"]["origem"].replace({"Cliente": "Indicação de cliente"}))
    if not a["atendidas"]:
        st.info("Nenhum trabalho fechado neste período. Os clientes aparecem aqui quando uma proposta é aprovada.")
    _kpis(
        _kpi("Clientes atendidas", str(a["atendidas"]),
             f'{a["novas"]} nova{"s" if a["novas"] != 1 else ""} · {a["voltaram"]} voltaram · '
             + _delta(a["atendidas"], c["atendidas"] or None, rot_comp)),
        _kpi("Valor por cliente", fmt_brl(a["valor_por_cliente"]) if a["valor_por_cliente"] else "—",
             _delta(a["valor_por_cliente"], c["valor_por_cliente"], rot_comp)),
        _kpi("Clientes que voltam", _pct(a["recompra"]),
             f'de {a["clientes_base"]} clientes, as que fecharam 2 ou mais trabalhos', "destaque"),
        _kpi("Dependência das 3 maiores", _pct(a["concentracao_top3"]), "do valor fechado no período"),
    )
    frases = []
    if a["atendidas"] and a["voltaram"] == 0 and a["novas"]:
        frases.append("Todo o período veio de clientes novas. Ofereça manutenção ou um segundo ambiente para quem já foi "
                      "atendida: vender para quem já confia em você custa menos do que conquistar uma cliente nova.")
    if a["concentracao_top3"] and a["concentracao_top3"] >= 0.6 and a["atendidas"] > 3:
        frases.append(f"As 3 maiores clientes respondem por {_pct(a['concentracao_top3'])} do valor: se uma delas parar, "
                      "o faturamento sente muito.")
    if not a["origem"].empty:
        o = a["origem"].iloc[0]
        frases.append(f"O canal que mais trouxe dinheiro foi {o['origem']}: {fmt_brl(o['valor'])} em {int(o['trabalhos'])} "
                      f"trabalho{'s' if o['trabalhos'] != 1 else ''}.")
    _insight(frases)
    c1, c2 = st.columns([1, 1], gap="large")
    with c1:
        _secao("De onde vêm as clientes")
        if a["origem"].empty:
            st.markdown('<div class="rl-vazio">Nada neste período.</div>', unsafe_allow_html=True)
        else:
            _tabela(a["origem"], moeda=["valor"], renomear={"origem": "Origem", "clientes": "Clientes",
                                                             "trabalhos": "Trabalhos", "valor": "Valor fechado"})
    with c2:
        _secao("Clientes que mais fecharam")
        if a["ranking"].empty:
            st.markdown('<div class="rl-vazio">Nada neste período.</div>', unsafe_allow_html=True)
        else:
            r = a["ranking"].assign(primeira_vez=a["ranking"]["primeira_vez"].map({True: "Nova", False: "Voltou"}))
            _tabela(r.head(15), moeda=["valor"], renomear={"cliente": "Cliente", "trabalhos": "Trabalhos",
                                                            "valor": "Valor fechado", "primeira_vez": ""})
    st.markdown('<div class="rl-def">Conta o valor das propostas aprovadas no período (data da aprovação). '
                '“Clientes que voltam” olha todo o histórico até o fim do período. A origem vem do cadastro da cliente.</div>',
                unsafe_allow_html=True)
    exp["Clientes"] = a["ranking"].rename(columns={"cliente": "Cliente", "trabalhos": "Trabalhos", "valor": "Valor fechado",
                                                    "primeira_vez": "Primeira vez no período"})
    exp["Origem das clientes"] = a["origem"].rename(columns={"origem": "Origem", "clientes": "Clientes",
                                                             "trabalhos": "Trabalhos", "valor": "Valor fechado"})


def _aba_propostas(propostas, ini, fim, cini, cfim, rot_comp, exp):
    a, c = rc.propostas_periodo(propostas, ini, fim), rc.propostas_periodo(propostas, cini, cfim)
    if not a["enviadas"]:
        st.info("Nenhuma proposta feita neste período.")
    _kpis(
        _kpi("Propostas feitas", str(a["enviadas"]),
             f'{fmt_brl(a["valor_enviado"])} · ' + _delta(a["enviadas"], c["enviadas"] or None, rot_comp)),
        _kpi("Taxa de fechamento", _pct(a["conversao"]),
             (_delta(a["conversao"], c["conversao"], rot_comp, pontos=True) if a["conversao"] is not None and c["conversao"] is not None
              else f'{a["fechadas"]} fechadas · {a["recusadas"]} recusadas'), "destaque"),
        _kpi("Fechamento em valor", _pct(a["conversao_valor"]),
             f'{fmt_brl(a["valor_fechado"])} fechados de {fmt_brl(a["valor_fechado"] + a["valor_recusado"])} respondidos'),
        _kpi("Tempo para fechar", "—" if a["dias_para_fechar"] is None else f'{a["dias_para_fechar"]:.0f} dias',
             "mediana, da proposta à aprovação"),
    )
    frases = []
    if a["ticket_fechado"] and a["ticket_recusado"] and a["ticket_recusado"] > a["ticket_fechado"] * 1.3:
        frases.append(f"As propostas recusadas custavam em média {fmt_brl(a['ticket_recusado'])}, contra "
                      f"{fmt_brl(a['ticket_fechado'])} das fechadas. Vale oferecer opções em etapas ou por ambiente.")
    if a["conversao"] is not None and a["conversao_valor"] is not None and a["conversao_valor"] < a["conversao"] - 0.15:
        frases.append("Você fecha mais propostas pequenas do que grandes: a taxa em valor está bem abaixo da taxa em quantidade.")
    if a["aguardando"]:
        frases.append(f"{a['aguardando']} proposta{'s' if a['aguardando'] != 1 else ''} deste período ainda "
                      f"aguarda{'m' if a['aguardando'] != 1 else ''} resposta ({fmt_brl(a['valor_aguardando'])}). "
                      "Veja a lista em Próximos meses.")
    _insight(frases)
    mens = rc.propostas_mensal(propostas, ini, fim)
    if len(mens) > 1 and mens["enviadas"].sum():
        _secao("Propostas por mês")
        fig = go.Figure()
        fig.add_bar(x=mens["mes"], y=mens["enviadas"], name="Feitas", marker_color="#D0D5DD")
        fig.add_bar(x=mens["mes"], y=mens["fechadas"], name="Fechadas", marker_color=NAVY)
        fig.add_bar(x=mens["mes"], y=mens["recusadas"], name="Recusadas", marker_color="#FDA29B")
        fig.update_layout(barmode="group", bargap=0.3)
        fig = _layout(fig, 280)
        fig.update_yaxes(tickprefix="")
        st.plotly_chart(fig, use_container_width=True, config={"displayModeBar": False})
    _secao("Por tipo de serviço")
    tipos = rc.propostas_por_tipo(propostas, ini, fim)
    if tipos.empty:
        st.markdown('<div class="rl-vazio">Nada neste período.</div>', unsafe_allow_html=True)
    else:
        _tabela(tipos, moeda=["valor_fechado", "ticket"], pct=["conversao"],
                renomear={"tipo": "Serviço", "enviadas": "Feitas", "fechadas": "Fechadas", "recusadas": "Recusadas",
                          "conversao": "Fechamento", "valor_fechado": "Valor fechado", "ticket": "Ticket médio"})
    st.markdown('<div class="rl-def">Considera as propostas feitas no período. Taxa de fechamento = fechadas ÷ (fechadas + recusadas): '
                'as que ainda aguardam resposta não entram na conta.</div>', unsafe_allow_html=True)
    exp["Propostas por tipo"] = tipos
    exp["Propostas por mês"] = mens


def _aba_trabalhos(mov, propostas, ini, fim, exp):
    t = rc.lucro_por_trabalho(mov, propostas, ini, fim)
    if t.empty:
        st.info("Nenhum trabalho fechado neste período.")
        return
    com_custo = t[t["custos"] > 0]
    receb = float(t["recebido"].sum())
    _kpis(
        _kpi("Trabalhos fechados", str(len(t)), f'{fmt_brl(t["valor"].sum())} contratados'),
        _kpi("Já recebido", fmt_brl(receb), f'{_pct(receb / t["valor"].sum() if t["valor"].sum() else None)} do contratado'),
        _kpi("Lucro dos trabalhos", fmt_brl(t["lucro"].sum()), "recebido − custos ligados ao trabalho", "destaque"),
        _kpi("Margem média", _pct(t["lucro"].sum() / receb if receb else None, 1),
             f'{len(com_custo)} de {len(t)} com custos lançados'),
    )
    frases = []
    baixa = t[(t["margem"].notna()) & (t["margem"] < 0.3)]
    if len(baixa):
        frases.append(f"{len(baixa)} trabalho{'s' if len(baixa) != 1 else ''} com margem abaixo de 30%. "
                      "Confira o preço desses serviços ou os custos com equipe e material.")
    if len(com_custo) < len(t):
        frases.append("Para o lucro por trabalho ficar certo, lance as despesas (assistente, material, fornecedor) "
                      "dentro da proposta no Financeiro.")
    _insight(frases)
    _tabela(t, moeda=["valor", "recebido", "custos", "lucro"], pct=["margem"],
            renomear={"numero": "Proposta", "cliente": "Cliente", "tipo": "Serviço", "valor": "Contratado",
                      "recebido": "Recebido", "custos": "Custos", "lucro": "Lucro", "margem": "Margem"})
    exp["Lucro por trabalho"] = t


def _aba_futuro(mov, propostas, hoje, exp):
    faixas, atrasados = rc.agenda_pendente(mov, hoje)
    # taxa de fechamento dos últimos 12 meses, para estimar quanto das propostas em aberto deve virar trabalho
    hist = rc.propostas_periodo(propostas, hoje - timedelta(days=365), hoje)
    ag = rc.propostas_aguardando(propostas, hoje)
    rec30 = float(faixas.loc[faixas["faixa"].isin(["Atrasado", "Até 30 dias"]), "a_receber"].sum())
    pag30 = float(faixas.loc[faixas["faixa"].isin(["Atrasado", "Até 30 dias"]), "a_pagar"].sum())
    v_ag = float(ag["valor"].sum()) if not ag.empty else 0.0
    esperado = v_ag * hist["conversao"] if hist["conversao"] is not None else None
    _kpis(
        _kpi("A receber em 30 dias", fmt_brl(rec30), "inclui os atrasados"),
        _kpi("A pagar em 30 dias", fmt_brl(pag30), "inclui os atrasados"),
        _kpi("Saldo dos próximos 30 dias", fmt_brl(rec30 - pag30), "a receber − a pagar", "destaque"),
        _kpi("Propostas aguardando", fmt_brl(v_ag),
             f'{len(ag)} em aberto' + (f' · devem virar ~{fmt_brl(esperado)}' if esperado else "")),
    )
    atr_rec = atrasados[atrasados["lado"] == "entra"] if not atrasados.empty else atrasados
    if not atr_rec.empty:
        st.markdown(f'<div class="rl-alerta"><b>{len(atr_rec)} recebimento{"s" if len(atr_rec) != 1 else ""} atrasado'
                    f'{"s" if len(atr_rec) != 1 else ""}</b>, somando {fmt_brl(atr_rec["valor"].sum())}. Cobre pelo WhatsApp '
                    'ou, se já recebeu, marque como Recebido no Financeiro.</div>', unsafe_allow_html=True)
    c1, c2 = st.columns([1, 1], gap="large")
    with c1:
        _secao("Agenda de recebimentos e pagamentos")
        f = faixas[(faixas["a_receber"] > 0) | (faixas["a_pagar"] > 0)].assign(saldo=lambda d: d["a_receber"] - d["a_pagar"])
        if f.empty:
            st.markdown('<div class="rl-vazio">Nada pendente.</div>', unsafe_allow_html=True)
        else:
            _tabela(f, moeda=["a_receber", "a_pagar", "saldo"],
                    renomear={"faixa": "Vencimento", "a_receber": "A receber", "a_pagar": "A pagar", "saldo": "Saldo"})
        if not atrasados.empty:
            _secao("Atrasados")
            d = atrasados.assign(tipo=atrasados["lado"].map({"entra": "Receber", "sai": "Pagar"}),
                                 venc=atrasados["d"].apply(lambda x: x.strftime("%d/%m/%Y") if x else "—"),
                                 dias=atrasados["d"].apply(lambda x: (hoje - x).days if x else None))
            _tabela(d, moeda=["valor"], renomear={"tipo": "", "descricao": "Descrição", "venc": "Venceu em",
                                                   "dias": "Dias", "valor": "Valor"})
    with c2:
        _secao("Propostas aguardando resposta")
        if ag.empty:
            st.markdown('<div class="rl-vazio">Nenhuma proposta em aberto.</div>', unsafe_allow_html=True)
        else:
            _tabela(ag, moeda=["valor"], renomear={"numero": "Proposta", "cliente": "Cliente", "valor": "Valor", "dias": "Dias parada"})
            paradas, velhas = ag[ag["dias"].fillna(0) > 7], ag[ag["dias"].fillna(0) > 60]
            if len(paradas):
                st.caption(f"{len(paradas)} sem resposta há mais de 7 dias: um lembrete simples costuma destravar a decisão."
                           + (f" {len(velhas)} estão paradas há mais de 60 dias e dificilmente vão fechar: marque-as como recusadas "
                              "nas Propostas para a taxa de fechamento e esta previsão ficarem certas." if len(velhas) else ""))
    st.markdown('<div class="rl-def">Pela data de vencimento das receitas e despesas pendentes. A estimativa das propostas usa a '
                'taxa de fechamento dos últimos 12 meses.</div>', unsafe_allow_html=True)
    exp["A receber e a pagar"] = faixas.rename(columns={"faixa": "Vencimento", "a_receber": "A receber", "a_pagar": "A pagar"})
    exp["Propostas aguardando"] = ag


# Excel já formatado: cabeçalho nas cores da marca, R$ e % nas colunas certas, larguras ajustadas,
# cabeçalho fixo e filtro em cada aba. As colunas são reconhecidas pelo nome.
_XL_NOMES = {"valor": "Valor", "valor_fechado": "Valor fechado", "ticket": "Ticket médio", "conversao": "Conversão",
             "margem": "Margem", "custos": "Custos", "lucro": "Lucro", "recebido": "Recebido", "dias": "Dias parada",
             "numero": "Proposta", "cliente": "Cliente", "tipo": "Tipo", "mes": "Mês", "enviadas": "Enviadas",
             "fechadas": "Fechadas", "recusadas": "Recusadas", "em_aberto": "Em aberto", "categoria": "Categoria"}
_XL_PCT = ("conversão", "margem", "taxa", "%")
_XL_INTEIRO = ("trabalhos", "clientes", "enviadas", "fechadas", "recusadas", "em aberto", "dias", "qtd")


def _xl_formato(coluna, serie):
    """Formato numérico do Excel pela coluna: %, inteiro ou R$ (o resto dos números é dinheiro)."""
    if not pd.api.types.is_numeric_dtype(serie) or pd.api.types.is_bool_dtype(serie):
        return None
    nome = str(coluna).lower()
    if any(k in nome for k in _XL_PCT):
        return "0.0%"
    if "proposta" in nome:
        return "0"
    if any(k in nome for k in _XL_INTEIRO):
        return "#,##0"
    return '"R$" #,##0.00;[Red]-"R$" #,##0.00'


def _excel(exp, titulo):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    navy, dourado, creme = "0D1B2A", "C9A84C", "F8F6F1"
    fino = Side(style="thin", color="E4E7EC")
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        try:
            from pages.perfil import carregar_perfil
            perfil = carregar_perfil() or {}
        except Exception:
            perfil = {}
        conta = perfil.get("empresa") or perfil.get("nome") or ""
        sobre = pd.DataFrame({"": ["Relatório", "Conta", "Gerado em", "Abas"],
                              " ": [titulo, conta, datetime.now().strftime("%d/%m/%Y %H:%M"),
                                    ", ".join(n for n, df in exp.items() if isinstance(df, pd.DataFrame) and not df.empty)]})
        sobre.to_excel(xw, sheet_name="Sobre", index=False, header=False, startrow=2)
        ws = xw.sheets["Sobre"]
        ws["A1"] = "Planner Organizer"
        ws["A1"].font = Font(bold=True, size=16, color=navy)
        for r in range(3, 7):
            ws.cell(r, 1).font = Font(bold=True, color="667085")
        ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 14, 60
        ws["B6"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.sheet_view.showGridLines = False

        for nome, df in exp.items():
            if not isinstance(df, pd.DataFrame) or df.empty:
                continue
            df = df.rename(columns=lambda c: _XL_NOMES.get(c, str(c).replace("_", " ").capitalize()
                                                            if "_" in str(c) else c))
            aba = nome[:31]
            df.to_excel(xw, sheet_name=aba, index=False, startrow=2)
            ws = xw.sheets[aba]
            ncol, nlin = len(df.columns), len(df) + 3
            ws["A1"] = nome
            ws["A1"].font = Font(bold=True, size=14, color=navy)
            ws["A2"] = titulo
            ws["A2"].font = Font(italic=True, size=9, color="667085")
            for j, col in enumerate(df.columns, 1):
                h = ws.cell(3, j)
                h.font, h.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=navy)
                h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                h.border = Border(bottom=Side(style="medium", color=dourado))
                fmt = _xl_formato(col, df[col])
                largura = len(str(col))
                for i in range(4, nlin + 1):
                    c = ws.cell(i, j)
                    if isinstance(c.value, str) and c.value.startswith("="):
                        c.data_type = "s"  # "= Lucro" é texto, não fórmula (antes aparecia #NAME?)
                    c.border = Border(bottom=fino)
                    if fmt:
                        c.number_format = fmt
                    elif hasattr(c.value, "year"):
                        c.number_format = "DD/MM/YYYY"
                    if i % 2 == 1:
                        c.fill = PatternFill("solid", fgColor=creme)
                    largura = max(largura, len(f"{c.value:,.2f}") + 4 if isinstance(c.value, (int, float)) and fmt
                                  else len(str(c.value or "")))
                ws.column_dimensions[get_column_letter(j)].width = min(max(largura + 2, 10), 50)
            ws.row_dimensions[3].height = 22
            ws.freeze_panes = "A4"
            ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{nlin}"
            ws.sheet_view.showGridLines = False
    return buf.getvalue()


def show():
    from utils.auth_guard import require_auth
    require_auth()
    st.markdown(CSS, unsafe_allow_html=True)
    st.markdown('<h1 class="page-title">Relatórios</h1>', unsafe_allow_html=True)
    st.markdown('<div class="rl-intro">Escolha um período e compare com outro. Para o mês de hoje, veja o Início; '
                'para lançar e conferir valores, o Financeiro.</div>', unsafe_allow_html=True)

    hoje = _hoje()
    c1, c2, c3 = st.columns([1, 1.4, 1])
    preset = c1.selectbox("Período", rc.PERIODOS, index=rc.PERIODOS.index("Últimos 3 meses"), key="rel_periodo")
    ini_c = fim_c = None
    if preset == "Personalizado":
        datas = c2.date_input("Datas", value=(hoje.replace(day=1), hoje), format="DD/MM/YYYY", key="rel_datas")
        if isinstance(datas, (list, tuple)) and len(datas) == 2:
            ini_c, fim_c = datas
        else:
            st.caption("Escolha a data inicial e a final.")
            return
    ini, fim = rc.intervalo(preset, hoje, ini_c, fim_c)
    comparar = c3.selectbox("Comparar com", ["Período anterior", "Mesmo período do ano passado"], key="rel_comp")
    cini, cfim = rc.periodo_anterior(ini, fim) if comparar == "Período anterior" else rc.mesmo_periodo_ano_passado(ini, fim)
    rot, rot_comp = rc.rotulo_periodo(ini, fim), rc.rotulo_periodo(cini, cfim)
    # período em andamento: compara até hoje com o mesmo número de dias do outro período
    ini, fim, cini, cfim, cortado = rc.ate_hoje(ini, fim, cini, cfim, hoje)
    if cortado:
        rot, rot_comp = f"{rot} até {fim:%d/%m}", f"{rot_comp} até {cfim:%d/%m}"
    if preset != "Personalizado":
        c2.markdown(f'<div style="padding-top:34px;color:#475467;font-size:.9rem"><b>{rot}</b> · comparado com {rot_comp}</div>',
                    unsafe_allow_html=True)

    try:
        mov, propostas, clientes = _carregar()
    except Exception as e:
        st.error(f"Não consegui carregar os dados: {e}")
        return

    exp = {}
    abas = st.tabs(["Resultado", "Clientes", "Propostas", "Trabalhos", "Próximos meses"])
    with abas[0]:
        _aba_resultado(mov, ini, fim, cini, cfim, rot, rot_comp, exp)
    with abas[1]:
        _aba_clientes(propostas, clientes, ini, fim, cini, cfim, rot_comp, exp)
    with abas[2]:
        _aba_propostas(propostas, ini, fim, cini, cfim, rot_comp, exp)
    with abas[3]:
        _aba_trabalhos(mov, propostas, ini, fim, exp)
    with abas[4]:
        _aba_futuro(mov, propostas, hoje, exp)

    st.divider()
    try:
        st.download_button("Baixar este relatório em Excel", data=_excel(exp, f"Relatório {rot} (comparado com {rot_comp})"),
                           icon=":material/table_view:",
                           file_name=f"relatorio_{ini:%Y-%m-%d}_{fim:%Y-%m-%d}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="rel_excel")
    except Exception as e:
        st.caption(f"Não foi possível gerar o Excel: {e}")
